"""Instrumentos de evaluación a partir de una GUÍA DE APRENDIZAJE ya generada.

La guía (JSON guardado junto al Word) trae la trazabilidad: fase, actividad del proyecto,
RAP, actividad de aprendizaje, evidencia y criterios de evaluación (del diseño curricular).
Con eso se arman:
- Lista de chequeo (desempeño / producto): indicadores observables ligados a los criterios.
- Rúbrica analítica de 4 niveles (producto): un criterio por fila, verbatim del diseño.
- Cuestionario (conocimiento): preguntas de selección múltiple con clave para el instructor.

Los instrumentos NO tienen formato institucional: se entregan en un Word propio con la línea
gráfica LogiLab SENA (verde #39A900). El juicio es SENA: Aprobado / No aprobado.
"""
from __future__ import annotations

import re
from pathlib import Path
import unicodedata

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERDE = "39A900"
VERDE_CLARO = "E8F5E0"
GRIS = "F2F2F2"
OSCURO = RGBColor(0x1F, 0x2D, 0x1A)

TIPOS = {
    "lista_chequeo": "Lista de chequeo",
    "rubrica": "Rúbrica analítica",
    "cuestionario": "Cuestionario",
}
NIVELES = [("excelente", "Excelente", 4), ("bueno", "Bueno", 3),
           ("aceptable", "Aceptable", 2), ("por_mejorar", "Por mejorar", 1)]
REGLA_DEFECTO = {
    "lista_chequeo": "Aprobado: cumple TODOS los indicadores. No aprobado: falta al menos uno.",
    "rubrica": "Aprobado: todos los criterios en nivel Aceptable o superior. No aprobado: algún criterio en «Por mejorar».",
    "cuestionario": "Aprobado: mínimo el 70 % de respuestas correctas. No aprobado: menos del 70 %.",
}


def _norm(t: str) -> str:
    t = unicodedata.normalize("NFD", str(t or "").lower())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


# ─────────────────────────────── evidencias de la guía ───────────────────────────────
_RE_TIPO = re.compile(r"evidencias?\s+de\s+(conocimiento|desempe[nñ]o|producto)\s*(?:\(([\d.]+)\))?\s*[:.\-–]?\s*", re.I)


def tipo_evidencia(texto: str) -> str:
    n = _norm(texto)
    if re.search(r"conocimiento|cuestionario|prueba escrita|examen|test", n):
        return "conocimiento"
    if re.search(r"producto|informe|documento|entregable|plan |formato|matriz|propuesta|ficha|presentacion", n):
        return "producto"
    return "desempeno"


def partir_evidencia(texto: str) -> list[tuple[str, str]]:
    """'Evidencia de conocimiento: X. Evidencia de producto: Y' → [(conocimiento, X), (producto, Y)].
    Las sub-evidencias numeradas ('Evidencia de conocimiento (3.3.1.1): X') conservan su número
    al inicio del texto: '(3.3.1.1) X'."""
    texto = str(texto or "").strip()
    marcas = list(_RE_TIPO.finditer(texto))
    if not marcas:
        return [(tipo_evidencia(texto), texto)] if texto else []
    salida = []
    for i, m in enumerate(marcas):
        fin = marcas[i + 1].start() if i + 1 < len(marcas) else len(texto)
        cuerpo = texto[m.end():fin].strip(" .;\n")
        if cuerpo and m.group(2):
            cuerpo = f"({m.group(2)}) {cuerpo}"
        if cuerpo:
            salida.append(("desempeno" if _norm(m.group(1)).startswith("desempe") else _norm(m.group(1)), cuerpo))
    return salida


def instrumento_sugerido(tipo: str) -> str:
    return {"conocimiento": "cuestionario", "producto": "rubrica"}.get(tipo, "lista_chequeo")


TECNICA = {"conocimiento": "formulación de preguntas", "desempeno": "observación sistemática",
           "producto": "valoración de producto"}
NOMBRE_EN_GUIA = {"cuestionario": "cuestionario", "lista_chequeo": "lista de chequeo", "rubrica": "rúbrica"}


def texto_guia(tipo: str, instrumento: str | None = None) -> str:
    """Texto de la columna «Técnicas e instrumentos» de la guía (fuente ÚNICA para guía e instrumentos)."""
    return (f"Técnica: {TECNICA.get(tipo, TECNICA['producto'])} · "
            f"Instrumento: {NOMBRE_EN_GUIA[instrumento or instrumento_sugerido(tipo)]}")


def instrumentos_de_evidencia(evidencia: str, elegidos: dict | None = None) -> str:
    """Columna 6 de la tabla 4 para una evidencia completa (puede traer conocimiento+desempeño+producto).
    `elegidos` = {texto_parte: instrumento} cuando el instructor cambió el sugerido."""
    partes = partir_evidencia(evidencia) or [("producto", evidencia)]
    return "\n".join(dict.fromkeys(texto_guia(t, (elegidos or {}).get(txt)) for t, txt in partes))


def _criterios_lista(c) -> list[str]:
    if isinstance(c, list):
        return [str(x).strip() for x in c if str(x).strip()]
    return [x.strip() for x in re.split(r"\n\s*\n|\n", str(c or "")) if x.strip()]


def evidencias_de_guia(guia: dict) -> list[dict]:
    """Lista plana de evidencias evaluables de la guía, cada una con su contexto."""
    base = guia.get("_evidencias")
    if not base:                                        # guías antiguas: desde la tabla 4
        base = []
        for fila in (guia.get("evidencias_tabla") or [])[1:]:
            if len(fila) >= 5:
                base.append({"fase": fila[0], "actividad_proyecto": fila[1], "rap": "",
                             "actividad": fila[2], "evidencia": fila[3], "criterios": fila[4]})
    items = []
    for n_aa, e in enumerate(base):
        for tipo, txt in partir_evidencia(e.get("evidencia", "")):
            sub = re.match(r"\((\d[\d.]*)\)\s*", txt)
            items.append({
                "id": len(items), "aa": n_aa, "integradora": bool(e.get("integradora")),
                "momento": sub.group(1) if sub else e.get("momento", ""), "tipo": tipo, "evidencia": txt,
                "saberes": ((e.get("sub_saberes") or {}).get(sub.group(1)) if sub else None) or e.get("saberes") or [],
                "fase": e.get("fase", ""), "actividad_proyecto": e.get("actividad_proyecto", ""),
                "rap": e.get("rap", ""), "actividad": e.get("actividad", ""),
                "criterios": _criterios_lista(e.get("criterios")),
                "instrumento": instrumento_sugerido(tipo),
            })
    return items


# ─────────────────────────────── contenido sin IA (plantilla) ───────────────────────────────
def _quitar_verbo_3p(c: str) -> str:
    c = c.rstrip(". ")
    if not c.isupper():
        return c
    # diseños en MAYÚSCULA → oración, conservando siglas (SST, NTC, ISO, EPP…)
    siglas = {"ISO", "NTC", "SENA", "PHVA", "EPP", "TIC", "DIAN", "RUT"}
    pal = [w if (w in siglas or (len(w) > 1 and not re.search(r"[AEIOUÁÉÍÓÚ]", w))) else w.lower() for w in c.split(" ")]
    t = " ".join(pal)
    return t[:1].upper() + t[1:]


def lista_base(item: dict) -> list[dict]:
    """Indicadores = criterios del diseño (verbatim), uno por fila."""
    crit = item.get("criterios") or [f"Entrega la evidencia: {item['evidencia']}"]
    return [{"indicador": _quitar_verbo_3p(c), "criterio": c} for c in crit]


def rubrica_base(item: dict) -> list[dict]:
    """Sin IA: criterio verbatim del diseño + descriptores graduados genéricos (cortos)."""
    return [{
        "criterio": c,
        "excelente": "Cumple el criterio de forma completa, precisa y argumentada; aporta valor adicional al proyecto.",
        "bueno": "Cumple el criterio de forma completa y correcta; quedan detalles menores por ajustar.",
        "aceptable": "Cumple lo mínimo exigido por el criterio, con apoyo del instructor.",
        "por_mejorar": "No evidencia el criterio; requiere plan de mejoramiento.",
    } for c in (item.get("criterios") or [item["evidencia"]])]


def cuestionario_base(item: dict, saberes: list[str], n: int = 5) -> list[dict]:
    """Sin IA solo se pueden proponer preguntas abiertas a partir de los saberes del diseño."""
    fuentes = [s for s in saberes if s] or item.get("criterios") or [item["evidencia"]]
    preguntas = []
    for s in fuentes[:n]:
        preguntas.append({"enunciado": f"Explica con tus palabras y con un ejemplo del proyecto: {s.rstrip('.')}.",
                          "criterio": "", "a": "", "b": "", "c": "", "d": "", "correcta": "",
                          "justificacion": "Respuesta abierta: valora la precisión conceptual y la aplicación al contexto."})
    return preguntas


def contenido_base(item: dict, instrumento: str, saberes: list[str]) -> list[dict]:
    if instrumento == "rubrica":
        return rubrica_base(item)
    if instrumento == "cuestionario":
        return cuestionario_base(item, item.get("saberes") or saberes)
    return lista_base(item)


# ─────────────────────────────── trazabilidad: nada por fuera ───────────────────────────────
def criterios_cubiertos(filas: list[dict]) -> set[str]:
    return {str(f.get("criterio", "")).strip() for f in filas or [] if str(f.get("criterio", "")).strip()}


def cobertura(items: list[dict], instrumentos: dict) -> dict:
    """`instrumentos` = {item_id: {"tipo", "contenido"}} (solo los elegidos).
    Revisa, por cada actividad de aprendizaje (AA) de la guía:
      - que TODAS sus evidencias tengan instrumento;
      - que TODOS sus criterios de evaluación queden evaluados en algún instrumento de esa AA;
      - que el instrumento coincida con lo que dice la guía (tabla 4)."""
    por_aa = {}
    for it in items:
        por_aa.setdefault(it["aa"], []).append(it)
    filas, sin_instr, sin_criterio, desalineados = [], [], [], []
    for aa, its in por_aa.items():
        crit = list(dict.fromkeys(c for it in its for c in it["criterios"]))
        cubiertos = set()
        for it in its:
            ins = instrumentos.get(it["id"])
            if not ins or not ins.get("contenido"):
                sin_instr.append(it)
                continue
            cubiertos |= criterios_cubiertos(ins["contenido"])
            if ins["tipo"] != it["instrumento"]:
                desalineados.append(it)
        faltan = [c for c in crit if c not in cubiertos]
        sin_criterio += [(its[0], c) for c in faltan]
        for it in its:
            ins = instrumentos.get(it["id"]) or {}
            mios = criterios_cubiertos(ins.get("contenido"))
            filas.append({
                "Momento": it.get("momento", ""), "AA": aa + 1, "Actividad de aprendizaje": it["actividad"],
                "Evidencia": f"{it['tipo'].replace('desempeno', 'desempeño').capitalize()}: {it['evidencia']}",
                "Instrumento": NOMBRE_EN_GUIA.get(ins.get("tipo"), "— sin instrumento —").capitalize(),
                "Criterios evaluados": ", ".join(f"CE{k + 1:02d}" for k, c in enumerate(crit) if c in mios) or "—",
                "Estado": "✅" if ins.get("contenido") and not faltan else "⚠️",
            })
    total_c = sum(len(dict.fromkeys(c for it in its for c in it["criterios"])) for its in por_aa.values())
    return {"filas": filas, "sin_instrumento": sin_instr, "sin_criterio": sin_criterio,
            "desalineados": desalineados, "total_criterios": total_c,
            "completo": not sin_instr and not sin_criterio}


def cerrar_cabos(items: list[dict], instrumentos: dict) -> dict:
    """Agrega a cada AA los criterios que ningún instrumento evalúa (en su primer instrumento
    de lista o rúbrica; si solo tiene cuestionario, una pregunta abierta ligada al criterio)."""
    rep = cobertura(items, instrumentos)
    for it, c in rep["sin_criterio"]:
        destino = next((x for x in items if x["aa"] == it["aa"] and x["id"] in instrumentos
                        and instrumentos[x["id"]].get("tipo") != "cuestionario"), None) or \
            next((x for x in items if x["aa"] == it["aa"] and x["id"] in instrumentos), None)
        if destino is None:
            continue
        ins = instrumentos[destino["id"]]
        fila = {"lista_chequeo": lista_base, "rubrica": rubrica_base}.get(ins["tipo"])
        if fila:
            ins["contenido"].append(fila({"criterios": [c], "evidencia": destino["evidencia"]})[0])
        else:
            ins["contenido"].append({"enunciado": f"Explica cómo se cumple en tu evidencia: {_quitar_verbo_3p(c)}.",
                                     "criterio": c, "a": "", "b": "", "c": "", "d": "", "correcta": "",
                                     "justificacion": "Respuesta abierta ligada al criterio de evaluación."})
    return instrumentos


def actualizar_guia(guia: dict, items: list[dict], instrumentos: dict) -> dict:
    """Devuelve una copia de la guía con la columna «Técnicas e instrumentos» (tabla 4) y el campo
    de instrumentos de 3.3/3.4 alineados con los instrumentos realmente construidos."""
    import copy
    g = copy.deepcopy(guia)
    elegidos = {}
    for it in items:
        if it["id"] in instrumentos:
            elegidos.setdefault(it["aa"], {})[it["evidencia"]] = instrumentos[it["id"]]["tipo"]
    evs = g.get("_evidencias") or []
    tabla = g.get("evidencias_tabla") or []
    for n, e in enumerate(evs):
        if n + 1 < len(tabla) and len(tabla[n + 1]) >= 6:
            tabla[n + 1][5] = instrumentos_de_evidencia(e.get("evidencia", ""), elegidos.get(n))
    for k in ("3.3", "3.4"):
        act = (g.get("actividades") or {}).get(k)
        aas = [n for n, e in enumerate(evs) if k in str(e.get("momento", ""))]
        if act is not None and aas:
            act["instrumentos"] = "\n".join(dict.fromkeys(
                instrumentos_de_evidencia(evs[n].get("evidencia", ""), elegidos.get(n)) for n in aas))
    return g


# ─────────────────────────────── Word (línea LogiLab SENA) ───────────────────────────────
def _sombrear(celda, hex_color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _bordes(tabla, color="9BBB8A"):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _anchos(tabla, anchos):
    """Anchos fijos (Word y LibreOffice): sin autoajuste, con tblGrid y ancho por celda."""
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for j, w in enumerate(anchos):
        tabla.columns[j].width = Cm(w)
        for c in tabla.columns[j].cells:
            c.width = Cm(w)


def _texto(celda, texto, negrita=False, tam=9, color=None, centro=False):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(texto or ""))
    r.bold = negrita; r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = color


def _tabla(doc, filas, anchos, encabezado=True, tam=9):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes(t)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            if encabezado and i == 0:
                _texto(c, v, True, tam, RGBColor(0xFF, 0xFF, 0xFF), True)
                _sombrear(c, VERDE)
            else:
                _texto(c, v, tam=tam)
    _anchos(t, anchos)
    return t


def _titulo(doc, texto, tam=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(texto); r.bold = True; r.font.size = Pt(tam)
    r.font.color.rgb = RGBColor(0x39, 0xA9, 0x00)
    return p


def _identificacion(doc, guia, item, nombre_inst):
    filas = [
        ("Programa de formación", f"{guia.get('programa', '')} {('(' + guia['codigo_programa'] + ')') if guia.get('codigo_programa') else ''}"),
        ("Proyecto formativo", guia.get("proyecto_formativo", "")),
        ("Fase / actividad del proyecto", f"{item['fase']} · {item['actividad_proyecto']}"),
        ("Competencia", guia.get("competencia", "")),
        ("Resultado de aprendizaje", item.get("rap", "") or "; ".join(guia.get("raps", []))),
        ("Actividad de aprendizaje", (f"[Guía {item['momento']}] " if item.get("momento") else "") + item.get("actividad", "")),
        ("Evidencia", f"{item['tipo'].replace('desempeno', 'desempeño').capitalize()}: {item['evidencia']}"),
        ("Instrumento", nombre_inst),
        ("Aprendiz / ficha", "______________________________   Ficha: ____________"),
        ("Fecha", "____ / ____ / ________"),
    ]
    t = doc.add_table(rows=len(filas), cols=2)
    _bordes(t)
    for i, (k, v) in enumerate(filas):
        _texto(t.cell(i, 0), k, True, 9); _sombrear(t.cell(i, 0), VERDE_CLARO)
        _texto(t.cell(i, 1), v, tam=9)
    _anchos(t, [5, 19.3])


def _juicio(doc, regla):
    _titulo(doc, "Juicio de evaluación", 10)
    t = doc.add_table(rows=2, cols=3)
    _bordes(t)
    for j, v in enumerate(("☐ APROBADO", "☐ NO APROBADO", "Observaciones / plan de mejoramiento")):
        _texto(t.cell(0, j), v, True, 9, centro=True); _sombrear(t.cell(0, j), VERDE_CLARO)
    _texto(t.cell(1, 2), "\n\n")
    _anchos(t, [4, 4, 16.3])
    p = doc.add_paragraph(); r = p.add_run("Regla: " + regla); r.italic = True; r.font.size = Pt(8)
    f = doc.add_table(rows=1, cols=2)
    _texto(f.cell(0, 0), "\n\n______________________________\nFirma del instructor", tam=9, centro=True)
    _texto(f.cell(0, 1), "\n\n______________________________\nFirma del aprendiz", tam=9, centro=True)
    _anchos(f, [12.15, 12.15])


LOGO = Path(__file__).resolve().parent.parent / "templates" / "logo_sena.png"
REGIONAL_CENTRO_DEFECTO = "Regional Guajira - Centro Industrial y de Energías Alternativas"


def regional_y_centro(texto: str) -> tuple[str, str]:
    """'Regional Guajira - Centro Industrial…' → ('Regional Guajira', 'Centro Industrial…')."""
    partes = [x.strip() for x in re.split(r"\s+[-–·|]\s+|\n", str(texto or "")) if x.strip()]
    reg = next((x for x in partes if _norm(x).startswith("regional")), "")
    cen = next((x for x in partes if _norm(x).startswith("centro")), "")
    resto = [x for x in partes if x not in (reg, cen)]
    if not reg and resto:
        reg = resto.pop(0)
    if not cen and resto:
        cen = resto.pop(0)
    return reg, cen


def _encabezado(doc, guia):
    """Encabezado en todas las páginas: logo SENA + regional + centro + programa."""
    reg, cen = regional_y_centro(guia.get("_regional_centro") or REGIONAL_CENTRO_DEFECTO)
    for s in doc.sections:
        s.header_distance = Cm(0.8)
        h = s.header
        h.paragraphs[0].text = ""
        t = h.add_table(rows=1, cols=2, width=Cm(24.3))
        _bordes(t, "FFFFFF")
        c0, c1 = t.cell(0, 0), t.cell(0, 1)
        if LOGO.exists():
            c0.paragraphs[0].add_run().add_picture(str(LOGO), height=Cm(1.6))
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        lineas = [("SERVICIO NACIONAL DE APRENDIZAJE · SENA", True, 10),
                  (" · ".join(x for x in (reg, cen) if x), True, 9),
                  (f"{guia.get('programa', '')}" + (f" · Código {guia['codigo_programa']}" if guia.get("codigo_programa") else ""), False, 8.5),
                  ("Instrumentos de evaluación · Sistema de Gestión de la Formación Profesional Integral", False, 8)]
        c1.text = ""
        for k, (txt, neg, tam) in enumerate(lineas):
            par = c1.paragraphs[0] if k == 0 else c1.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            r = par.add_run(txt); r.bold = neg; r.font.size = Pt(tam)
            if k < 2:
                r.font.color.rgb = RGBColor(0x39, 0xA9, 0x00) if k == 0 else OSCURO
        _anchos(t, [2.6, 21.7])
        linea = h.add_paragraph()
        pPr = linea._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "12"); b.set(qn("w:color"), VERDE)
        bdr.append(b); pPr.append(bdr)


def _pie(doc):
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("LogiLab SENA · Instrumentos de evaluación · By Carlos Barrios · Instructor")
        r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x39, 0xA9, 0x00)


def generar_instrumentos_docx(guia: dict, instrumentos: list[dict], ruta: str,
                              matriz: list[dict] | None = None) -> str:
    """`instrumentos` = [{"item": evidencia, "tipo": lista_chequeo|rubrica|cuestionario,
    "contenido": [...], "regla": str}]. Un instrumento por página; la clave del
    cuestionario va al final, en hoja aparte, solo para el instructor."""
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.top_margin = Cm(3.4)
    _encabezado(doc, guia)
    _pie(doc)

    claves = []
    if matriz:
        p = doc.add_paragraph(); r = p.add_run("MATRIZ DE TRAZABILIDAD · GUÍA → EVIDENCIAS → INSTRUMENTOS")
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x39, 0xA9, 0x00)
        doc.add_paragraph(f"Guía: {guia.get('competencia', '')} · {guia.get('fase_proyecto', '')}").runs[0].font.size = Pt(9)
        cols = ["Momento", "AA", "Actividad de aprendizaje", "Evidencia", "Instrumento", "Criterios evaluados", "Estado"]
        filas = [["Momento", "AA", "Actividad de aprendizaje", "Evidencia", "Instrumento", "Criterios", "✓"]]
        filas += [[str(m.get(c, "")) for c in cols] for m in matriz]
        _tabla(doc, filas, [1.6, 1, 7.5, 6.6, 2.6, 3.8, 1.2], tam=8)
        nota = doc.add_paragraph("Cada actividad de aprendizaje de la guía tiene todas sus evidencias con instrumento y "
                                 "todos sus criterios de evaluación (CE, del diseño curricular) evaluados.")
        nota.runs[0].italic = True; nota.runs[0].font.size = Pt(8)
    for n, ins in enumerate(instrumentos):
        if n or matriz:
            doc.add_page_break()
        item, tipo, cont = ins["item"], ins["tipo"], ins.get("contenido") or []
        nombre = TIPOS.get(tipo, tipo)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(f"INSTRUMENTO N.° {n + 1} · {nombre.upper()}"); r.bold = True; r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0x39, 0xA9, 0x00)
        _identificacion(doc, guia, item, nombre)

        if tipo == "lista_chequeo":
            _titulo(doc, "Indicadores de evaluación", 10)
            filas = [["N.°", "Indicador", "Sí", "No", "Observaciones"]]
            filas += [[str(i + 1), c.get("indicador", ""), "☐", "☐", ""] for i, c in enumerate(cont)]
            _tabla(doc, filas, [1, 15, 1.2, 1.2, 5.9])
        elif tipo == "rubrica":
            _titulo(doc, "Rúbrica analítica (4 niveles)", 10)
            filas = [["Criterio de evaluación"] + [f"{t} ({p})" for _, t, p in NIVELES] + ["Puntaje"]]
            filas += [[c.get("criterio", "")] + [c.get(k, "") for k, _, _ in NIVELES] + [""] for c in cont]
            _tabla(doc, filas, [5, 4.45, 4.45, 4.45, 4.45, 1.5], tam=8)
        else:
            _titulo(doc, "Cuestionario", 10)
            p = doc.add_paragraph("Instrucción: lee cada enunciado y marca con una X la opción correcta. "
                                  "En las preguntas abiertas responde con tus palabras.")
            p.runs[0].font.size = Pt(9)
            for i, q in enumerate(cont):
                pq = doc.add_paragraph(); pq.paragraph_format.space_after = Pt(2)
                rq = pq.add_run(f"{i + 1}. {q.get('enunciado', '')}"); rq.bold = True; rq.font.size = Pt(10)
                ops = [(k, q.get(k, "")) for k in "abcd" if str(q.get(k, "")).strip()]
                if ops:
                    for k, v in ops:
                        po = doc.add_paragraph(f"     ☐ {k}) {v}"); po.paragraph_format.space_after = Pt(0)
                        po.runs[0].font.size = Pt(9.5)
                else:
                    doc.add_paragraph("_" * 120 + "\n" + "_" * 120).runs[0].font.size = Pt(9)
            claves.append((item, cont))
        _juicio(doc, ins.get("regla") or REGLA_DEFECTO.get(tipo, ""))

    for item, cont in claves:
        doc.add_page_break()
        p = doc.add_paragraph(); r = p.add_run("CLAVE DE RESPUESTAS · SOLO PARA EL INSTRUCTOR")
        r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x39, 0xA9, 0x00)
        doc.add_paragraph(f"Evidencia: {item['evidencia']}").runs[0].font.size = Pt(9)
        crit = item.get("criterios") or []
        ce = lambda c: f"CE{crit.index(c) + 1:02d}" if c in crit else "—"
        filas = [["N.°", "Respuesta correcta", "Criterio", "Justificación"]]
        filas += [[str(i + 1), (q.get("correcta", "") or "Abierta").upper(), ce(q.get("criterio", "")),
                   q.get("justificacion", "")] for i, q in enumerate(cont)]
        _tabla(doc, filas, [1, 3.3, 1.8, 18.2])
    doc.save(ruta)
    return ruta
